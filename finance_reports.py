# -*- coding: utf-8 -*-
"""
Web3 企业级财务 ERP - 报表生成引擎
"""
import logging
from datetime import datetime
from typing import Dict, List, Any, Optional
from io import BytesIO
import aiosqlite

logger = logging.getLogger(__name__)

try:
    from weasyprint import WeasyPrint
    WEASYPRINT_AVAILABLE = True
except ImportError:
    WEASYPRINT_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

async def collect_finance_summary(month: str = None) -> Dict[str, Any]:
    """收集财务摘要数据"""
    from config import DB_NAME
    
    if not month:
        month = datetime.now().strftime("%Y-%m")
    
    year, month_num = month.split("-")
    start_date = f"{year}-{month_num}-01"
    
    if int(month_num) == 12:
        end_date = f"{int(year)+1}-01-01"
    else:
        end_date = f"{year}-{int(month_num)+1:02d}-01"
    
    async with aiosqlite.connect(DB_NAME) as db:
        db.row_factory = aiosqlite.Row
        
        cursor = await db.execute('''
            SELECT SUM(base_fiat_value) as total
            FROM finance_ledgers
            WHERE direction = 'in' 
            AND created_at >= ? AND created_at < ?
        ''', (start_date, end_date))
        inflow = (await cursor.fetchone())["total"] or 0
        
        cursor = await db.execute('''
            SELECT SUM(base_fiat_value) as total
            FROM finance_ledgers
            WHERE direction = 'out' 
            AND created_at >= ? AND created_at < ?
        ''', (start_date, end_date))
        outflow = (await cursor.fetchone())["total"] or 0
        
        cursor = await db.execute('''
            SELECT a.code, a.name, a.type,
                   SUM(CASE WHEN l.direction = 'in' THEN l.base_fiat_value ELSE 0 END) as inflow,
                   SUM(CASE WHEN l.direction = 'out' THEN l.base_fiat_value ELSE 0 END) as outflow,
                   COUNT(*) as tx_count
            FROM finance_ledgers l
            JOIN chart_of_accounts a ON l.account_id = a.account_id
            WHERE l.created_at >= ? AND l.created_at < ?
            GROUP BY a.account_id
            ORDER BY a.code
        ''', (start_date, end_date))
        
        categories = []
        for row in await cursor.fetchall():
            categories.append({
                "code": row["code"],
                "name": row["name"],
                "type": row["type"],
                "inflow": row["inflow"] or 0,
                "outflow": row["outflow"] or 0,
                "tx_count": row["tx_count"]
            })
        
        return {
            "period": month,
            "total_inflow": inflow,
            "total_outflow": outflow,
            "net_flow": inflow - outflow,
            "categories": categories,
            "generated_at": datetime.now().isoformat()
        }

async def collect_ledger_details(month: str = None, limit: int = 100) -> List[Dict[str, Any]]:
    """收集总账明细"""
    from config import DB_NAME
    
    if not month:
        month = datetime.now().strftime("%Y-%m")
    
    year, month_num = month.split("-")
    start_date = f"{year}-{month_num}-01"
    
    if int(month_num) == 12:
        end_date = f"{int(year)+1}-01-01"
    else:
        end_date = f"{year}-{int(month_num)+1:02d}-01"
    
    async with aiosqlite.connect(DB_NAME) as db:
        db.row_factory = aiosqlite.Row
        cursor = await db.execute('''
            SELECT l.ledger_id, l.direction, l.amount, l.token_symbol,
                   l.base_fiat_value, l.counterparty, l.description,
                   l.reconciled_status, l.created_at,
                   a.code, a.name
            FROM finance_ledgers l
            JOIN chart_of_accounts a ON l.account_id = a.account_id
            WHERE l.created_at >= ? AND l.created_at < ?
            ORDER BY l.created_at DESC
            LIMIT ?
        ''', (start_date, end_date, limit))
        
        return [dict(row) for row in await cursor.fetchall()]

def generate_pdf_report(summary: Dict[str, Any], details: List[Dict[str, Any]]) -> Optional[bytes]:
    """生成企业级 PDF 报表"""
    if not WEASYPRINT_AVAILABLE:
        logger.error("WeasyPrint 未安装")
        return None
    
    html_content = _build_pdf_html(summary, details)
    
    try:
        pdf_bytes = WeasyPrint(string=html_content).write_pdf()
        return pdf_bytes
    except Exception as e:
        logger.error(f"PDF 生成失败: {e}")
        return None

def _build_pdf_html(summary: Dict[str, Any], details: List[Dict[str, Any]]) -> str:
    """构建 PDF HTML 模板"""
    
    period = summary["period"]
    inflow = summary["total_inflow"]
    outflow = summary["total_outflow"]
    net_flow = summary["net_flow"]
    
    category_rows = ""
    for cat in summary["categories"]:
        category_rows += f"<tr><td>{cat['code']}</td><td>{cat['name']}</td><td>{cat['type']}</td><td style='text-align:right'>${cat['inflow']:.2f}</td><td style='text-align:right'>${cat['outflow']:.2f}</td><td style='text-align:right'>{cat['tx_count']}</td></tr>"
    
    detail_rows = ""
    for detail in details[:50]:
        detail_rows += f"<tr><td>{detail['created_at'][:10]}</td><td>{detail['code']}</td><td>{detail['direction'].upper()}</td><td>{detail['amount']} {detail['token_symbol']}</td><td style='text-align:right'>${detail['base_fiat_value']:.2f}</td><td>{detail['counterparty'] or '-'}</td><td>{detail['reconciled_status']}</td></tr>"
    
    html = f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><style>
body{{font-family:Arial,sans-serif;color:#333;margin:0;padding:40px}}
.header{{border-bottom:2px solid #2c3e50;padding-bottom:20px;margin-bottom:30px}}
.header h1{{font-size:24px;color:#2c3e50}}
.meta{{text-align:right;font-size:12px;color:#7f8c8d}}
.summary{{display:grid;grid-template-columns:1fr 1fr 1fr;gap:20px;margin-bottom:30px}}
.summary-card{{background:#ecf0f1;padding:15px;border-radius:4px;border-left:4px solid #3498db}}
.summary-card.positive{{border-left-color:#27ae60}}
.summary-card.negative{{border-left-color:#e74c3c}}
.summary-card .label{{font-size:12px;color:#7f8c8d}}
.summary-card .value{{font-size:20px;font-weight:bold;color:#2c3e50}}
table{{width:100%;border-collapse:collapse;margin-bottom:30px;font-size:11px}}
table th{{background:#34495e;color:white;padding:10px;text-align:left}}
table td{{padding:8px 10px;border-bottom:1px solid #ecf0f1}}
table tr:nth-child(even){{background:#f8f9fa}}
.section-title{{font-size:14px;font-weight:bold;color:#2c3e50;margin-top:20px;margin-bottom:10px;border-bottom:1px solid #bdc3c7;padding-bottom:5px}}
.footer{{margin-top:40px;padding-top:20px;border-top:1px solid #bdc3c7;font-size:10px;color:#7f8c8d;display:flex;justify-content:space-between}}
</style></head><body>
<div class="header"><h1>💼 企业财务报表</h1><p style="font-size:12px;color:#7f8c8d">Web3 加密资产管理系统</p><div class="meta"><p>报告期: {period}</p><p>生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p></div></div>
<div class="summary">
<div class="summary-card positive"><div class="label">总入金</div><div class="value">${inflow:,.2f}</div></div>
<div class="summary-card negative"><div class="label">总出金</div><div class="value">${outflow:,.2f}</div></div>
<div class="summary-card {'positive' if net_flow >= 0 else 'negative'}"><div class="label">净流量</div><div class="value">${net_flow:,.2f}</div></div>
</div>
<div class="section-title">📊 分类汇总</div>
<table><thead><tr><th>科目代码</th><th>科目名称</th><th>类型</th><th style="text-align:right">入金</th><th style="text-align:right">出金</th><th style="text-align:right">笔数</th></tr></thead><tbody>{category_rows}</tbody></table>
<div class="section-title">📝 交易明细（前50条）</div>
<table><thead><tr><th>日期</th><th>科目</th><th>方向</th><th>金额</th><th style="text-align:right">法币价值</th><th>交易对手</th><th>对账状态</th></tr></thead><tbody>{detail_rows}</tbody></table>
<div class="footer"><div><p>制表人: _______________</p></div><div><p>审核人: _______________</p></div><div><p>批准人: _______________</p></div></div>
</body></html>"""
    
    return html

def generate_excel_report(summary: Dict[str, Any], details: List[Dict[str, Any]]) -> Optional[bytes]:
    """生成多 Sheet 的 Excel 报表"""
    if not OPENPYXL_AVAILABLE:
        logger.error("openpyxl 未安装")
        return None
    
    try:
        wb = Workbook()
        wb.remove(wb.active)
        
        _add_summary_sheet(wb, summary)
        _add_category_sheet(wb, summary)
        _add_details_sheet(wb, details)
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    
    except Exception as e:
        logger.error(f"Excel 生成失败: {e}")
        return None

def _add_summary_sheet(wb, summary: Dict[str, Any]):
    """添加摘要 Sheet"""
    ws = wb.create_sheet("财务摘要", 0)
    ws['A1'] = "💼 企业财务摘要"
    ws['A1'].font = Font(size=16, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    ws.merge_cells('A1:D1')
    ws['A3'] = "报告期:"
    ws['B3'] = summary["period"]
    ws['A4'] = "生成时间:"
    ws['B4'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    ws['A6'] = "指标"
    ws['B6'] = "金额 (USD)"
    for cell in ['A6', 'B6']:
        ws[cell].font = Font(bold=True, color="FFFFFF")
        ws[cell].fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    ws['A7'] = "总入金"
    ws['B7'] = summary["total_inflow"]
    ws['B7'].number_format = '$#,##0.00'
    ws['A8'] = "总出金"
    ws['B8'] = summary["total_outflow"]
    ws['B8'].number_format = '$#,##0.00'
    ws['A9'] = "净流量"
    ws['B9'] = summary["net_flow"]
    ws['B9'].number_format = '$#,##0.00'
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 20

def _add_category_sheet(wb, summary: Dict[str, Any]):
    """添加分类汇总 Sheet"""
    ws = wb.create_sheet("分类汇总", 1)
    ws['A1'] = "📊 分类汇总"
    ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    ws.merge_cells('A1:F1')
    headers = ["科目代码", "科目名称", "类型", "入金 (USD)", "出金 (USD)", "笔数"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    for row, cat in enumerate(summary["categories"], 4):
        ws.cell(row=row, column=1).value = cat["code"]
        ws.cell(row=row, column=2).value = cat["name"]
        ws.cell(row=row, column=3).value = cat["type"]
        ws.cell(row=row, column=4).value = cat["inflow"]
        ws.cell(row=row, column=4).number_format = '$#,##0.00'
        ws.cell(row=row, column=5).value = cat["outflow"]
        ws.cell(row=row, column=5).number_format = '$#,##0.00'
        ws.cell(row=row, column=6).value = cat["tx_count"]
    for col in range(1, 7):
        ws.column_dimensions[chr(64+col)].width = 18

def _add_details_sheet(wb, details: List[Dict[str, Any]]):
    """添加明细 Sheet"""
    ws = wb.create_sheet("交易明细", 2)
    ws['A1'] = "📝 交易明细"
    ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    ws.merge_cells('A1:G1')
    headers = ["日期", "科目", "方向", "金额", "法币价值", "交易对手", "对账状态"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    for row, detail in enumerate(details, 4):
        ws.cell(row=row, column=1).value = detail['created_at'][:10]
        ws.cell(row=row, column=2).value = detail['code']
        ws.cell(row=row, column=3).value = detail['direction'].upper()
        ws.cell(row=row, column=4).value = f"{detail['amount']} {detail['token_symbol']}"
        ws.cell(row=row, column=5).value = detail['base_fiat_value']
        ws.cell(row=row, column=5).number_format = '$#,##0.00'
        ws.cell(row=row, column=6).value = detail['counterparty'] or '-'
        ws.cell(row=row, column=7).value = detail['reconciled_status']
    for col in range(1, 8):
        ws.column_dimensions[chr(64+col)].width = 18

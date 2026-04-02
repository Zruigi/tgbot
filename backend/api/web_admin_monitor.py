# -*- coding: utf-8 -*-
"""Web 管理端监控与潜在用户相关接口模块。"""
import re
import logging
from aiohttp import web
import database

logger = logging.getLogger(__name__)

async def api_list_monitors(request):
    """列出所有钱包监听（安全加固：过滤当前管理员的监听，防止跨租户数据泄露）"""
    try:
        active_only = request.query.get('active_only', 'false').lower() == 'true'
        rows = await database.get_all_monitors_for_admin(active_only=active_only)

        # 修复 High: 监听管理无所有权验证 —— 只返回属于当前管理员的监听记录
        admin_id = config.ADMIN_ID
        rows = [r for r in rows if int(r.get('user_id', 0)) == admin_id]

        # 附加近 20 笔出入金统计，便于前端展示风险与概览
        for row in rows:
            txs = await database.get_recent_monitor_transactions(row['monitor_id'], limit=20)
            row['inflow_20'] = round(sum(t['amount'] for t in txs if t.get('direction') == 'in'), 2)
            row['outflow_20'] = round(sum(t['amount'] for t in txs if t.get('direction') == 'out'), 2)
            row['profit_20'] = round(row['inflow_20'] - row['outflow_20'], 2)

        return web.json_response({"status": "ok", "data": rows})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


def validate_wallet_address(address: str, network: str) -> bool:
    """
    校验钱包地址格式是否符合对应链的规范。
    修复 P0: 地址输入无格式验证（CVSS 9.1）——防止注入恶意字符串。
    """
    if not address or len(address) > 100:
        return False
    net = network.lower()
    if net in ('trc20', 'tron'):
        return bool(_TRON_ADDR_RE.match(address))
    if net in ('bsc', 'eth', 'erc20'):
        return bool(_ETH_BSC_ADDR_RE.match(address))
    return False


async def api_create_monitor(request):
    """新增钱包监听（管理端）（安全加固：地址格式验证）"""
    try:
        data = await request.json()
        user_id = int(data.get('user_id') or 0)
        address = (data.get('address') or '').strip()
        network = (data.get('network') or '').strip().lower()
        remark = (data.get('remark') or data.get('tag') or '').strip()

        if not user_id or not address or network not in ('trc20', 'bsc'):
            return web.json_response({"status": "error", "message": "参数不完整或网络类型错误"}, status=400)

        # 修复 P0: 地址输入无格式验证
        if not validate_wallet_address(address, network):
            logger.warning(f"非法地址格式拦截: address={address!r} network={network}")
            return web.json_response({"status": "error", "message": "钱包地址格式不合法"}, status=400)

        await database.save_wallet_monitor(user_id, address, network, remark=remark)
        await notify_web_clients("monitor_update", {"action": "created", "user_id": user_id, "address": address, "network": network.upper(), "remark": remark})
        return web.json_response({"status": "ok"})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_delete_monitor(request):
    """删除钱包监听（管理端）"""
    try:
        monitor_id = int(request.match_info['monitor_id'])
        ok = await database.delete_monitor_admin(monitor_id)
        if ok:
            await notify_web_clients("monitor_update", {"action": "deleted", "monitor_id": monitor_id})
            return web.json_response({"status": "ok"})
        return web.json_response({"status": "error", "message": "监听不存在"}, status=404)
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_update_monitor_remark(request):
    """更新监听备注"""
    try:
        monitor_id = int(request.match_info['monitor_id'])
        data = await request.json()
        remark = (data.get('remark') or '').strip()
        ok = await database.update_monitor_remark(monitor_id, remark)
        if ok:
            await notify_web_clients('monitor_update', {'action': 'remark_updated', 'monitor_id': monitor_id, 'remark': remark})
            return web.json_response({'status': 'ok'})
        return web.json_response({'status': 'error', 'message': '监听不存在'}, status=404)
    except Exception as e:
        return web.json_response({'status': 'error', 'message': str(e)}, status=500)


async def api_monitor_pnl_summary(request):
    """获取当日/周/月盈亏台账汇总（支持单地址、多地址过滤）"""
    try:
        uid = request.query.get('user_id')
        user_id = int(uid) if uid else None

        addresses_raw = (request.query.get('addresses') or '').strip()
        addresses = [x.strip() for x in addresses_raw.split(',') if x.strip()] if addresses_raw else []

        monitor_ids_raw = (request.query.get('monitor_ids') or '').strip()
        monitor_ids = [int(x) for x in monitor_ids_raw.split(',') if x.strip().isdigit()] if monitor_ids_raw else []

        day = await database.get_monitor_pnl_summary('day', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids)
        day_prev = await database.get_monitor_pnl_summary('day', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids, previous=True)
        week = await database.get_monitor_pnl_summary('week', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids)
        week_prev = await database.get_monitor_pnl_summary('week', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids, previous=True)
        month = await database.get_monitor_pnl_summary('month', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids)
        month_prev = await database.get_monitor_pnl_summary('month', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids, previous=True)

        def calc_mom(curr: float, prev: float) -> Optional[float]:
            if prev == 0:
                return None if curr == 0 else 100.0
            return ((curr - prev) / abs(prev)) * 100.0

        series_day = await database.get_monitor_period_series('day', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids)
        series_week = await database.get_monitor_period_series('week', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids)
        series_month = await database.get_monitor_period_series('month', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids)
        tx_list = await database.get_monitor_transactions_list('month', user_id=user_id, addresses=addresses, monitor_ids=monitor_ids, limit=300)

        return web.json_response({
            'status': 'ok',
            'data': {
                'day': day,
                'week': week,
                'month': month,
                'comparison': {
                    'day_vs_prev_day': {
                        'current': day['profit'],
                        'previous': day_prev['profit'],
                        'change': day['profit'] - day_prev['profit'],
                        'change_percent': calc_mom(day['profit'], day_prev['profit']),
                    },
                    'week_vs_prev_week': {
                        'current': week['profit'],
                        'previous': week_prev['profit'],
                        'change': week['profit'] - week_prev['profit'],
                        'change_percent': calc_mom(week['profit'], week_prev['profit']),
                    },
                    'month_vs_prev_month': {
                        'current': month['profit'],
                        'previous': month_prev['profit'],
                        'change': month['profit'] - month_prev['profit'],
                        'change_percent': calc_mom(month['profit'], month_prev['profit']),
                    },
                },
                'series': {
                    'day': series_day,
                    'week': series_week,
                    'month': series_month,
                },
                'transactions': tx_list,
            }
        })
    except Exception as e:
        return web.json_response({'status': 'error', 'message': str(e)}, status=500)


async def api_update_monitor_tx_meta(request):
    """更新链上流水业务字段"""
    try:
        data = await request.json()
        tx_id = (data.get('tx_id') or '').strip()
        if not tx_id:
            return web.json_response({'status': 'error', 'message': 'tx_id 不能为空'}, status=400)

        ok = await database.update_monitor_transaction_meta(
            tx_id=tx_id,
            income_reason=(data.get('income_reason') or '').strip() or None,
            expense_reason=(data.get('expense_reason') or '').strip() or None,
            payer=(data.get('payer') or '').strip() or None,
            payee=(data.get('payee') or '').strip() or None,
        )
        if not ok:
            return web.json_response({'status': 'error', 'message': '流水不存在'}, status=404)
        return web.json_response({'status': 'ok'})
    except Exception as e:
        return web.json_response({'status': 'error', 'message': str(e)}, status=500)


async def api_batch_update_monitor_tx_meta(request):
    """批量更新链上流水业务字段"""
    try:
        data = await request.json()
        items = data.get('items') or []
        if not isinstance(items, list) or len(items) == 0:
            return web.json_response({'status': 'error', 'message': 'items 不能为空'}, status=400)

        updated = await database.batch_update_monitor_transaction_meta(items)
        return web.json_response({'status': 'ok', 'updated': updated})
    except Exception as e:
        return web.json_response({'status': 'error', 'message': str(e)}, status=500)


async def api_export_monitor_transactions(request):
    """导出链上流水 Excel（可按地址/多个地址导出）"""
    try:
        period = (request.query.get('period') or 'month').strip().lower()
        if period not in ('day', 'week', 'month'):
            period = 'month'

        uid = request.query.get('user_id')
        user_id = int(uid) if uid else None

        addresses_raw = (request.query.get('addresses') or '').strip()
        addresses = [x.strip() for x in addresses_raw.split(',') if x.strip()] if addresses_raw else []

        monitor_ids_raw = (request.query.get('monitor_ids') or '').strip()
        monitor_ids = [int(x) for x in monitor_ids_raw.split(',') if x.strip().isdigit()] if monitor_ids_raw else []

        annotated_only = (request.query.get('annotated_only') or 'false').lower() == 'true'

        pnl = await database.get_monitor_pnl_summary(period, user_id=user_id, addresses=addresses, monitor_ids=monitor_ids)
        rows = await database.get_monitor_transactions_for_export(
            start_ms=pnl['start_ms'],
            end_ms=pnl['end_ms'],
            user_id=user_id,
            addresses=addresses,
            monitor_ids=monitor_ids,
            annotated_only=annotated_only,
        )

        try:
            from openpyxl import Workbook
        except Exception:
            return web.json_response({'status': 'error', 'message': '缺少 openpyxl 依赖，请先安装'}, status=500)

        wb = Workbook()
        ws = wb.active
        ws.title = '链上流水'
        ws.append([
            '时间', '地址', '网络', '方向', '金额(USDT)', '对方地址',
            '收入原因', '支出原因', '谁打款', '打款给谁', '风险等级', '风险备注', 'TXID'
        ])

        for r in rows:
            dt = datetime.fromtimestamp(int(r['tx_timestamp']) / 1000).strftime('%Y-%m-%d %H:%M:%S')
            ws.append([
                dt,
                r.get('address') or '',
                (r.get('network') or '').upper(),
                '收入' if r.get('direction') == 'in' else '支出',
                float(r.get('amount') or 0),
                r.get('counterparty') or '',
                r.get('income_reason') or '',
                r.get('expense_reason') or '',
                r.get('payer') or '',
                r.get('payee') or '',
                r.get('risk_level') or '',
                r.get('risk_note') or '',
                r.get('tx_id') or '',
            ])

        summary = wb.create_sheet('汇总')
        summary.append(['周期', period])
        summary.append(['开始时间', datetime.fromtimestamp(pnl['start_ms'] / 1000).strftime('%Y-%m-%d %H:%M:%S')])
        summary.append(['结束时间', datetime.fromtimestamp(pnl['end_ms'] / 1000).strftime('%Y-%m-%d %H:%M:%S')])
        summary.append(['收入', round(float(pnl['inflow']), 2)])
        summary.append(['支出', round(float(pnl['outflow']), 2)])
        summary.append(['盈亏', round(float(pnl['profit']), 2)])

        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)

        filename = f"monitor_export_{period}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        headers = {
            'Content-Disposition': f'attachment; filename="{filename}"'
        }
        return web.Response(
            body=bio.getvalue(),
            headers=headers,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        return web.json_response({'status': 'error', 'message': str(e)}, status=500)


async def _sync_potential_users_from_group_messages(limit: int = 2000):
    """从已记录的群组消息中回填潜在用户。"""
    from handlers.group_monitor import DEFAULT_KEYWORDS, _profile_tags_from_keywords

    limit = max(1, min(int(limit), 10000))

    async with aiosqlite.connect(config.DB_NAME) as db:
        db.row_factory = aiosqlite.Row

        async with db.execute(
            """
            SELECT group_id, group_title, custom_keywords
            FROM tracked_groups
            WHERE is_monitored = 1
            """
        ) as cursor:
            groups = await cursor.fetchall()

        if not groups:
            return {
                "scanned_messages": 0,
                "matched_messages": 0,
                "upserted_users": 0,
                "groups": 0,
            }

        group_kw_map = {}
        group_title_map = {}
        group_ids = []
        for g in groups:
            gid = int(g["group_id"])
            group_ids.append(gid)
            group_title_map[gid] = g["group_title"] or ""
            custom_kws = [k.strip() for k in (g["custom_keywords"] or "").split(",") if k.strip()]
            group_kw_map[gid] = list(set(DEFAULT_KEYWORDS + custom_kws))

        placeholders = ",".join(["?"] * len(group_ids))
        query = f"""
            SELECT group_id, user_id, username, content
            FROM group_messages
            WHERE group_id IN ({placeholders})
              AND sender_type = 'user'
              AND content IS NOT NULL
              AND TRIM(content) <> ''
            ORDER BY COALESCE(message_ts, CAST(strftime('%s', timestamp) AS INTEGER)) DESC
            LIMIT ?
        """

        params = tuple(group_ids + [limit])
        async with db.execute(query, params) as cursor:
            rows = await cursor.fetchall()

    scanned_messages = len(rows)
    matched_messages = 0
    upserted_users = set()

    for row in rows:
        group_id = int(row["group_id"])
        text = (row["content"] or "").strip()
        if not text:
            continue

        keywords_pool = group_kw_map.get(group_id, DEFAULT_KEYWORDS)
        text_lower = text.lower()
        matched_keywords = [kw for kw in keywords_pool if kw and kw.lower() in text_lower]
        profile_tags = _profile_tags_from_keywords(matched_keywords, text)

        if not matched_keywords and not profile_tags:
            continue

        matched_messages += 1
        user_id = int(row["user_id"])
        await database.upsert_potential_user(
            user_id=user_id,
            username=row["username"] or f"User{user_id}",
            first_name=None,
            last_name=None,
            message_text=text,
            keywords=",".join(sorted(set(matched_keywords))),
            profile_tags=profile_tags,
            group_id=group_id,
            group_title=group_title_map.get(group_id, ""),
        )
        upserted_users.add(user_id)

    return {
        "scanned_messages": scanned_messages,
        "matched_messages": matched_messages,
        "upserted_users": len(upserted_users),
        "groups": len(group_ids),
    }


async def api_sync_potential_users(request):
    """手动触发潜在用户回填同步。"""
    try:
        limit = int(request.query.get('limit', 2000))
        stats = await _sync_potential_users_from_group_messages(limit=limit)

        global _tags_cache
        _tags_cache = {"data": [], "expire_at": 0}

        await notify_web_clients("potential_sync", {"stats": stats})

        summary = (
            f"同步完成：扫描 {stats['scanned_messages']} 条消息，"
            f"命中 {stats['matched_messages']} 条，"
            f"更新 {stats['upserted_users']} 位潜在用户。"
        )
        return web.json_response({"status": "ok", "summary": summary, "data": stats})
    except Exception as e:
        logger.error(f"同步潜在用户失败: {e}")
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_get_potential_users(request):
    """获取潜在用户列表 (支持分页与画像/关键词过滤)"""
    try:
        limit = min(int(request.query.get('limit', 50)), 500)
        offset = int(request.query.get('offset', 0))
        keyword = request.query.get('keyword')
        tag = request.query.get('tag')
        group_id = request.query.get('group_id')
        if group_id: group_id = int(group_id)
        
        result = await database.get_potential_users(limit=limit, offset=offset, keyword=keyword, tag=tag, group_id=group_id)
        return web.json_response({"status": "ok", "data": result['users'], "total": result['total']})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_get_potential_tags(request):
    """获取所有潜在用户画像标签 (60s 缓存)"""
    global _tags_cache
    try:
        now = time.time()
        if now > _tags_cache["expire_at"]:
            tags = await database.get_potential_tags()
            _tags_cache = {"data": tags, "expire_at": now + 60}
        return web.json_response({"status": "ok", "data": _tags_cache["data"], "total": len(_tags_cache["data"])})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_list_groups(request):
    """获取所有追踪的群组列表 (分页 2026 规格，含实时状态)"""
    try:
        page = int(request.query.get('page', 1)); size = int(request.query.get('size', 20))
        result = await database.get_tracked_groups_paginated(page, size)
        
        # 增加实时状态标记
        from datetime import datetime, timedelta
        now = datetime.now()
        threshold = now - timedelta(minutes=10)
        
        for g in result['groups']:
            updated_at = datetime.fromisoformat(g['updated_at'].replace(' ', 'T'))
            g['is_online'] = updated_at > threshold
            
        return web.json_response({"status": "ok", "data": result['groups'], "total": result['total']})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_get_monitor_stats(request):
    """2026 规格：获取群组实时监测面板数据"""
    try:
        groups = await database.get_tracked_groups_with_stats()
        online_count = sum(1 for g in groups if g['is_online'])
        monitored_count = sum(1 for g in groups if g['is_monitored'])
        
        return web.json_response({
            "status": "ok",
            "summary": {
                "total": len(groups),
                "online": online_count,
                "monitored": monitored_count
            },
            "groups": groups
        })
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_toggle_group_listen(request):
    """切换群组监听状态"""
    try:
        data = await request.json()
        group_id = int(data.get('group_id')); listen_status = bool(data.get('listen_status'))
        await database.set_group_listen_status(group_id, listen_status)
        from handlers.group_monitor import clear_group_cache
        clear_group_cache(group_id)
        await notify_web_clients("group_listen_update", {"group_id": group_id, "new_status": listen_status})
        return web.json_response({"code": 200, "msg": "状态已更新", "new_status": listen_status})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_toggle_group_revoke_sync(request):
    """切换群组信息校对（撤回同步删除）开关"""
    try:
        data = await request.json()
        group_id = int(data.get('group_id'))
        enabled = bool(data.get('enabled'))
        await database.set_group_delete_on_revoke(group_id, enabled)
        await notify_web_clients("group_revoke_sync_update", {"group_id": group_id, "enabled": enabled})
        return web.json_response({"status": "ok", "enabled": enabled})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_create_revoke_sync_chat(request):
    """保存信息校对暂存群 ID（管理员手动建群拉入 bot 后提交）"""
    try:
        if not _bot_instance:
            return web.json_response({"status": "error", "message": "bot 未初始化"}, status=500)
        data = await request.json()
        chat_id_input = str(data.get('chat_id', '')).strip()
        if not chat_id_input:
            # 未传 chat_id：返回当前已保存的暂存群信息
            existing = await database.get_admin_setting('revoke_sync_chat_id')
            if existing:
                try:
                    chat = await _bot_instance.get_chat(int(existing))
                    return web.json_response({"status": "ok", "chat_id": int(existing), "title": chat.title, "already_exists": True})
                except Exception:
                    return web.json_response({"status": "ok", "chat_id": int(existing), "title": "(获取失败)", "already_exists": True})
            return web.json_response({"status": "ok", "chat_id": None, "already_exists": False})
        # 验证 bot 是否在该群内
        try:
            chat_id = int(chat_id_input)
            chat = await _bot_instance.get_chat(chat_id)
        except Exception as e:
            return web.json_response({"status": "error", "message": f"无法获取该群信息，请确认 bot 已在群内且群 ID 正确：{e}"}, status=400)
        await database.save_admin_setting('revoke_sync_chat_id', str(chat_id))
        logger.info(f"[信息校对] 暂存群已设置: {chat_id} ({chat.title})")
        return web.json_response({"status": "ok", "chat_id": chat_id, "title": chat.title, "already_exists": False})
    except Exception as e:
        logger.error(f"api_create_revoke_sync_chat error: {e}")
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_get_revoke_sync_chat(request):
    """获取当前校对暂存群信息"""
    try:
        existing = await database.get_admin_setting('revoke_sync_chat_id')
        if not existing:
            return web.json_response({"status": "ok", "chat_id": None})
        if not _bot_instance:
            return web.json_response({"status": "ok", "chat_id": int(existing)})
        try:
            chat = await _bot_instance.get_chat(int(existing))
            return web.json_response({"status": "ok", "chat_id": int(existing), "title": chat.title})
        except Exception:
            return web.json_response({"status": "ok", "chat_id": int(existing), "title": "(群信息获取失败)"})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_promote_potential_user(request):
    """潜在用户转正（安全加固：校验用户是否在潜在用户列表中，防止任意 user_id 提升）"""
    try:
        user_id = int(request.match_info['user_id'])

        # 修复 High: 潜在用户提升无校验 —— 校验 user_id 是否在潜在用户表中
        async with aiosqlite.connect(config.DB_NAME, timeout=10) as db:
            cursor = await db.execute(
                "SELECT 1 FROM potential_users WHERE user_id = ?",
                (user_id,)
            )
            row = await cursor.fetchone()
            if not row:
                logger.warning(f"权限拦截: 尝试提升不在潜在用户列表中的 user_id={user_id}")
                return web.json_response({"status": "error", "message": "该用户不在潜在用户列表中"}, status=404)

        ok = await database.promote_potential_user_to_member(user_id)
        if ok:
            await notify_web_clients("potential_promoted", {"user_id": user_id})
        return web.json_response({"status": "ok" if ok else "error"})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_mark_read(request):
    """标记已读"""
    try:
        user_id = int(request.match_info['user_id'])
        await database.mark_messages_as_read(user_id)
        return web.json_response({"status": "ok"})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


async def api_delete_chat(request):
    """软删除会话"""
    try:
        user_id = int(request.match_info['user_id'])
        await database.soft_delete_chat(user_id)
        return web.json_response({"status": "ok"})
    except Exception as e:
        return web.json_response({"status": "error", "message": str(e)}, status=500)


